from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "nazarudin_TPLE004_Pertemuan2.docx"


def set_font(run, name="Times New Roman", size=12, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.5):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_keep_with_next(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    ppr.append(keep)


def add_title(doc, text, size):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(paragraph, after=4, line=1.0)
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=True)
    return paragraph


def add_identity_line(doc, label, value):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(1.55))
    set_paragraph_spacing(paragraph, after=2, line=1.0)
    label_run = paragraph.add_run(label)
    set_font(label_run, size=12.5, bold=True)
    separator = paragraph.add_run("\t: ")
    set_font(separator, size=12.5, bold=True)
    value_run = paragraph.add_run(value)
    set_font(value_run, size=12.5, bold=True)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    if level == 1:
        set_paragraph_spacing(paragraph, before=12, after=7, line=1.0)
        run = paragraph.add_run(text)
        set_font(run, size=14, bold=True)
    else:
        set_paragraph_spacing(paragraph, before=8, after=4, line=1.0)
        run = paragraph.add_run(text)
        set_font(run, size=12.5, bold=True)
    return paragraph


def add_body(doc, text, justify=True):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(paragraph, after=6, line=1.5)
    run = paragraph.add_run(text)
    set_font(run)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(paragraph, after=3, line=1.35)
    run = paragraph.add_run(text)
    set_font(run)
    return paragraph


def add_number(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(paragraph, after=3, line=1.35)
    run = paragraph.add_run(text)
    set_font(run)
    return paragraph


def add_tree(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.08)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    set_font(run, size=12, bold=True)
    return paragraph


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)

    if "Template Subtitle" not in [style.name for style in doc.styles]:
        subtitle = doc.styles.add_style("Template Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    else:
        subtitle = doc.styles["Template Subtitle"]
    subtitle.font.name = "Times New Roman"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    subtitle.font.size = Pt(22)
    subtitle.font.bold = True
    subtitle.font.color.rgb = RGBColor(0, 0, 0)


def build_document():
    doc = Document()
    configure_styles(doc)
    props = doc.core_properties
    props.author = "Nazarudin"
    props.title = "Laporan Tugas Pertemuan 2"
    props.subject = "JFrame Form Java 8 NetBeans 8.2"
    props.comments = "Template layout disesuaikan dengan project latihan2"

    add_title(doc, "LAPORAN TUGAS PERTEMUAN 2", 24)
    add_title(doc, "MATA KULIAH PEMROGRAMAN 2", 22)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    add_identity_line(doc, "NAMA", "Nazarudin")
    add_identity_line(doc, "NIM", "231011450485")
    add_identity_line(doc, "KELAS", "TPLE004")
    add_identity_line(doc, "MATAKUL", "Pemrograman 2")
    add_identity_line(doc, "PROJECT", "latihan2")

    add_heading(doc, "1. Tujuan Praktikum")
    add_number(doc, "Memahami pembuatan project Java menggunakan NetBeans 8.2.")
    add_number(doc, "Memahami hubungan antara project, source package, class Java, dan file desain JFrame Form.")
    add_number(doc, "Membuat antarmuka pengguna grafis sederhana menggunakan komponen Java Swing.")
    add_number(doc, "Menerapkan event tombol untuk menampilkan dan membersihkan output pada form.")

    add_heading(doc, "2. Dasar Teori")
    add_heading(doc, "2.1 Pengenalan NetBeans 8.2", level=2)
    add_body(
        doc,
        "NetBeans 8.2 adalah Integrated Development Environment berbasis Java yang menyediakan editor kode, "
        "compiler, debugger, dan GUI Builder. GUI Builder membantu pembuatan antarmuka desktop dengan cara "
        "menempatkan komponen visual pada form. Project ini menggunakan Java 8 agar sesuai dengan lingkungan "
        "praktikum dan konfigurasi NetBeans 8.2."
    )
    add_heading(doc, "2.2 Struktur Project Java di NetBeans", level=2)
    add_body(
        doc,
        "Project Java di NetBeans memiliki struktur yang memisahkan source code, konfigurasi project, dan hasil "
        "build. Folder src menyimpan package dan file program. Folder nbproject menyimpan metadata project "
        "berbasis Ant, sedangkan build dan dist digunakan saat project dikompilasi atau dibuat menjadi JAR."
    )
    add_heading(doc, "2.3 Java Swing dan JFrame Form", level=2)
    add_body(
        doc,
        "Java Swing merupakan toolkit bawaan Java untuk membuat aplikasi desktop. JFrame berfungsi sebagai "
        "jendela utama aplikasi. Pada JFrame Form ini digunakan JLabel untuk teks, JTextField untuk input nama, "
        "JButton untuk menjalankan aksi, dan JPanel sebagai wadah komponen. NetBeans menyimpan desain GUI pada "
        "file .form dan kode inisialisasinya pada metode initComponents()."
    )

    add_heading(doc, "3. Alat dan Bahan")
    add_bullet(doc, "Laptop atau PC dengan sistem operasi Windows.")
    add_bullet(doc, "Java Development Kit (JDK) versi 8.")
    add_bullet(doc, "NetBeans IDE 8.2.")
    add_bullet(doc, "Project Java berbasis Ant dengan nama latihan2.")

    add_heading(doc, "4. Langkah Kerja")
    add_number(doc, "Membuka aplikasi NetBeans IDE 8.2.")
    add_number(doc, "Membuat atau membuka project Java Application dengan nama latihan2.")
    add_number(doc, "Memastikan compiler source dan target project diatur ke Java 1.8.")
    add_number(doc, "Membuat package utama bernama nazarudin_pertemuan2 pada Source Packages.")
    add_number(doc, "Membuat JFrame Form pada package tersebut dengan nama class nazarudin_Latihan2.")
    add_number(doc, "Mendesain form dengan judul, petunjuk, input Nama, tombol Tampilkan, tombol Bersihkan, dan label output.")
    add_number(doc, "Menambahkan event tombol: Tampilkan memberi sapaan sesuai nama, sedangkan Bersihkan mengembalikan kondisi awal form.")
    add_number(doc, "Menjalankan program dengan menekan Shift + F6 atau tombol Run File.")

    add_heading(doc, "5. Hasil dan Pembahasan")
    add_heading(doc, "5.1 Struktur Akhir Project", level=2)
    add_body(doc, "Setelah langkah-langkah dilakukan, struktur project yang terbentuk adalah sebagai berikut:", justify=False)
    add_tree(
        doc,
        "latihan2\n"
        "|-- Source Packages\n"
        "|   \\-- nazarudin_pertemuan2\n"
        "|       |-- nazarudin_Latihan2.java\n"
        "|       \\-- nazarudin_Latihan2.form\n"
        "|-- Libraries\n"
        "|-- nbproject\n"
        "|   |-- build-impl.xml\n"
        "|   |-- project.properties\n"
        "|   \\-- project.xml\n"
        "|-- build.xml\n"
        "\\-- manifest.mf"
    )
    add_body(
        doc,
        "File nazarudin_Latihan2.java berisi class JFrame Form, komponen Swing, layout GroupLayout, serta event "
        "handler. File nazarudin_Latihan2.form berisi informasi desain yang digunakan oleh GUI Builder NetBeans."
    )

    add_heading(doc, "5.2 Komponen JFrame Form", level=2)
    add_bullet(doc, "JLabel: judul form, petunjuk, label Nama, dan informasi output.")
    add_bullet(doc, "JTextField: tempat pengguna memasukkan nama.")
    add_bullet(doc, "JButton Tampilkan: menampilkan sapaan pada label output.")
    add_bullet(doc, "JButton Bersihkan: menghapus isi input dan mengembalikan pesan output awal.")
    add_bullet(doc, "GroupLayout: mengatur posisi komponen agar tersusun rapi.")

    add_heading(doc, "5.3 Hasil Output", level=2)
    add_body(doc, "Saat program dijalankan, akan muncul jendela JFrame dengan judul Latihan 2 - JFrame Form. "
             "Di dalamnya terdapat input nama dan dua tombol aksi.")
    add_number(doc, "Jika pengguna mengisi nama Nazarudin lalu menekan Tampilkan, output menjadi: Halo, Nazarudin! Selamat belajar JFrame Form.")
    add_number(doc, "Jika input kosong, program menampilkan pesan informasi agar nama diisi terlebih dahulu.")
    add_number(doc, "Jika Bersihkan ditekan, JTextField dikosongkan dan output kembali menjadi Output akan tampil di sini.")

    add_heading(doc, "6. Kesimpulan")
    add_body(doc, "Berdasarkan praktikum yang telah dilakukan, dapat disimpulkan bahwa:")
    add_number(doc, "NetBeans 8.2 mempermudah pembuatan aplikasi desktop Java melalui fitur GUI Builder.")
    add_number(doc, "Struktur project latihan2 memisahkan konfigurasi NetBeans, source package, class Java, dan file desain form.")
    add_number(doc, "JFrame Form nazarudin_Latihan2 berhasil menampilkan antarmuka sederhana dan merespons aksi pengguna melalui tombol.")

    closing = doc.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(closing, before=10, after=0, line=1.5)
    run = closing.add_run("Terima Kasih.")
    set_font(run, size=12)

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
