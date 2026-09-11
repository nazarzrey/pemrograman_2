from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "Dokumentasi_Latihan2_Nazarudin.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def format_table(table, header=True):
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_borders(cell)
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
                    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
                    run.font.size = Pt(9.5)
        if header and row_index == 0:
            set_repeat_table_header(row)
            for cell in row.cells:
                set_cell_shading(cell, "1F4E79")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
        elif row_index % 2 == 0:
            for cell in row.cells:
                set_cell_shading(cell, "F5F8FC")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for cell, text in zip(hdr, headers):
        cell.text = text
    for row_values in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row_values):
            cell.text = text
    set_table_widths(table, widths)
    format_table(table, header=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def set_run_font(run, name="Arial", size=10.5, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_body_paragraph(doc, text="", bold_prefix=None):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.12
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, bold=True)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_number(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_code_block(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.16)
    paragraph.paragraph_format.right_indent = Inches(0.16)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.0
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement("w:" + edge)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "4")
        node.set(qn("w:color"), "D9D9D9")
        borders.append(node)
    p_pr.append(borders)
    run = paragraph.add_run(text)
    set_run_font(run, name="Consolas", size=9)
    return paragraph


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)

    for style_name, size in (("Heading 1", 14), ("Heading 2", 11.5)):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    if "Subtitle Custom" not in [style.name for style in doc.styles]:
        subtitle = doc.styles.add_style("Subtitle Custom", WD_STYLE_TYPE.PARAGRAPH)
    else:
        subtitle = doc.styles["Subtitle Custom"]
    subtitle.font.name = "Arial"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    subtitle.font.size = Pt(11.5)
    subtitle.font.color.rgb = RGBColor(0, 0, 0)
    subtitle.paragraph_format.space_after = Pt(14)


def build_document():
    doc = Document()
    style_document(doc)
    props = doc.core_properties
    props.author = "Nazarudin"
    props.title = "Dokumentasi Project Latihan 2"
    props.subject = "JFrame Form Java 8 NetBeans 8.2"
    props.comments = "Dokumentasi tugas Latihan 2"

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    title.add_run("Dokumentasi Project Latihan 2")

    subtitle = doc.add_paragraph(style="Subtitle Custom")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Pembuatan JFrame Form Java 8 Dengan NetBeans 8.2")

    intro = doc.add_paragraph(style="Normal")
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.paragraph_format.space_after = Pt(14)
    intro.paragraph_format.line_spacing = 1.1
    run = intro.add_run(
        "Dokumen ini menjelaskan hasil pembuatan project latihan JFrame Form "
        "sederhana, susunan file project, komponen antarmuka, dan alur outputnya."
    )
    set_run_font(run, size=10.5)

    doc.add_heading("Identitas Mahasiswa", level=1)
    add_table(
        doc,
        ["Keterangan", "Data"],
        [
            ("Nama", "Nazarudin"),
            ("NIM", "231011450485"),
            ("Kelas", "TPLE004"),
            ("Mata Kuliah", "Pemrograman 2"),
            ("Tanggal Dokumentasi", "12 September 2026"),
        ],
        [Inches(1.75), Inches(4.75)],
    )

    doc.add_heading("Identitas dan Tujuan Project", level=1)
    add_body_paragraph(
        doc,
        "Project ini diberi nama latihan2 dan dibuat menggunakan Java 8 dengan IDE NetBeans 8.2. "
        "Package utama menggunakan pola nama mahasiswa dan pertemuan, yaitu "
        "nazarudin_pertemuan2. JFrame Form digunakan untuk menampilkan antarmuka desktop sederhana "
        "berbasis Java Swing."
    )
    add_bullet(doc, "Mengenal pembuatan JFrame Form melalui GUI Builder NetBeans.")
    add_bullet(doc, "Menggunakan komponen JLabel, JTextField, JButton, dan JPanel.")
    add_bullet(doc, "Mengatur tampilan komponen dengan GroupLayout.")
    add_bullet(doc, "Membuat event tombol untuk menampilkan dan membersihkan output.")

    doc.add_heading("Struktur Project", level=1)
    add_body_paragraph(
        doc,
        "Struktur berikut menunjukkan hubungan antara file konfigurasi NetBeans, source code Java, "
        "dan file desain JFrame Form."
    )
    add_code_block(
        doc,
        "latihan2/\n"
        "|-- build.xml\n"
        "|-- manifest.mf\n"
        "|-- nbproject/\n"
        "|   |-- build-impl.xml\n"
        "|   |-- project.properties\n"
        "|   \\-- project.xml\n"
        "|-- src/\n"
        "|   \\-- nazarudin_pertemuan2/\n"
        "|       |-- nazarudin_Latihan2.java\n"
        "|       \\-- nazarudin_Latihan2.form\n"
        "|-- DOKUMENTASI_STRUKTUR_PROJECT.md\n"
        "|-- WORKLOG.md\n"
        "\\-- Dokumentasi_Latihan2_Nazarudin.docx"
    )
    add_table(
        doc,
        ["Bagian", "Fungsi"],
        [
            ("build.xml", "Entry point perintah Ant untuk Clean and Build project."),
            ("manifest.mf", "Menentukan class utama ketika project dibuat menjadi file JAR."),
            ("nbproject", "Menyimpan metadata dan konfigurasi project NetBeans berbasis Ant."),
            ("src", "Folder source code Java dan package project."),
            ("nazarudin_Latihan2.java", "Class JFrame Form, layout, komponen Swing, dan event tombol."),
            ("nazarudin_Latihan2.form", "Metadata desain yang dibaca GUI Builder NetBeans."),
            ("WORKLOG.md", "Catatan proses pembuatan dan status verifikasi project."),
        ],
        [Inches(2.1), Inches(4.4)],
    )

    doc.add_heading("Desain JFrame Form", level=1)
    add_body_paragraph(
        doc,
        "Saat dijalankan, form memiliki judul Latihan 2 - JFrame Form dan menampilkan komponen berikut."
    )
    add_table(
        doc,
        ["Komponen", "Keterangan"],
        [
            ("JLabel", "Menampilkan judul, petunjuk, label Nama, dan area output."),
            ("JTextField", "Tempat pengguna memasukkan nama."),
            ("Tombol Tampilkan", "Memproses nama dan menampilkan sapaan pada area output."),
            ("Tombol Bersihkan", "Mengosongkan input dan mengembalikan teks output awal."),
            ("GroupLayout", "Mengatur posisi komponen secara terstruktur pada JPanel."),
        ],
        [Inches(2.1), Inches(4.4)],
    )

    doc.add_heading("Alur Output", level=1)
    add_number(doc, "Jalankan class nazarudin_Latihan2.")
    add_number(doc, "Ketik nama pada JTextField.")
    add_number(doc, "Klik tombol Tampilkan.")
    add_number(doc, "Form menampilkan sapaan, misalnya: Halo, Nazarudin! Selamat belajar JFrame Form.")
    add_number(doc, "Klik Bersihkan untuk mengosongkan form dan mengembalikan output awal.")
    add_body_paragraph(
        doc,
        "Jika tombol Tampilkan ditekan ketika input masih kosong, program menampilkan pesan informasi "
        "agar pengguna mengisi nama terlebih dahulu."
    )

    doc.add_heading("Cara Membuka Project di NetBeans 8.2", level=1)
    add_number(doc, "Buka NetBeans 8.2.")
    add_number(doc, "Pilih File > Open Project.")
    add_number(doc, "Pilih folder project yang berisi nbproject dan build.xml.")
    add_number(doc, "Pastikan Java Platform yang digunakan adalah JDK 8.")
    add_number(doc, "Klik kanan project, pilih Clean and Build, kemudian jalankan project.")

    doc.add_heading("Kesimpulan", level=1)
    add_body_paragraph(
        doc,
        "Project latihan2 berhasil disusun sebagai project Java berbasis Ant dengan JFrame Form "
        "sederhana. Package nazarudin_pertemuan2 berisi class nazarudin_Latihan2 dan file desain "
        "GUI Builder. Form dapat menerima input nama, menampilkan output sapaan, dan membersihkan "
        "kembali isi form."
    )

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
